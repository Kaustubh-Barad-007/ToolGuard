import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_demo_document():
    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_run = title_p.add_run("🛡️ ToolGuard: Live Demonstration Script & Pitch Guide")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(22)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(18, 18, 18)

    # Subtitle
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    sub_run = sub_p.add_run("Complete 3-Minute Demo Walkthrough for Hackathon Judges, Investors & Technical Audiences")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(11.5)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(100, 116, 139)

    def add_heading_styled(text, level=1):
        h = doc.add_heading(text, level=level)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        for r in h.runs:
            r.font.name = "Calibri"
            if level == 1:
                r.font.size = Pt(15)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0, 196, 117)  # ToolGuard Green
            elif level == 2:
                r.font.size = Pt(13)
                r.font.bold = True
                r.font.color.rgb = RGBColor(30, 41, 59)
        return h

    # Section 1: Executive Overview & The 10-Second Pitch
    add_heading_styled("1. Executive Overview & The 10-Second Pitch", 1)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(
        "\"Traditional security tools only scan for known vulnerabilities (CVEs) in packages. "
        "They are completely blind to Capability Drift — when a trusted developer tool, npm script, or AI agent "
        "silently gains admin rights or network exfiltration access. ToolGuard brings Zero-Trust Capability Security "
        "to developer tools by cryptographically freezing permissions and alerting in real time before unauthorized code executes.\""
    )
    r.font.italic = True
    r.font.size = Pt(10.5)

    # Section 2: Preparation Checklist
    add_heading_styled("2. Pre-Demo Setup Checklist (Do This 2 Minutes Before)", 1)
    items = [
        ("Live Web Console URL: ", "Open https://toolguard-app.vercel.app in your browser."),
        ("Local IDE (Optional for Real-Time Sync): ", "Open VS Code / Cursor in any repo with the ToolGuard extension installed, OR open a terminal and run `toolguard serve`."),
        ("Terminal Tab: ", "Have a terminal window open in your project directory ready for quick commands."),
        ("1-Click Fallback Mode: ", "The web dashboard has a 1-click '⚡ Load Active IDE Project' button if running without an active local terminal.")
    ]
    for label, desc in items:
        bp = doc.add_paragraph(style="List Bullet")
        bp.paragraph_format.space_after = Pt(4)
        r1 = bp.add_run(label)
        r1.font.bold = True
        bp.add_run(desc)

    # Section 3: The 4-Step Interactive Live Demo
    add_heading_styled("3. The 4-Step Live Demo Walkthrough (3 Minutes Total)", 1)

    steps = [
        (
            "STEP 1: The Trust Anchor — Freeze Baseline (0:00 - 0:45)",
            "Establish the cryptographic baseline of authorized workspace tools.",
            "toolguard quickstart   (or click '⚡ Load Active IDE Project' on Web Dashboard)",
            "\"Every modern project runs npm scripts, MCP server tools, and background tasks. On Day 1, ToolGuard takes a tamper-proof cryptographic snapshot of these tools using canonical SHA-256 fingerprinting. Notice how dev-runner, npm:dev, and MCP filesystem are frozen into baseline.json. The status is 100% green and verified.\""
        ),
        (
            "STEP 2: The Attack — Simulate Capability Drift (0:45 - 1:30)",
            "Inject an unauthorized permission escalation or exfiltration endpoint.",
            "toolguard threat-test   (or click 'Simulate Incident' in Web Console / Settings)",
            "\"Now, imagine a compromised dependency or prompt injection alters our script. Notice what happened: npm:dev secretly gained admin rights and an external network endpoint to exfiltrate data. Traditional CVE scanners see zero alerts because the package version didn't change! But look at ToolGuard: within 50 milliseconds, our VS Code status bar turns RED, the CLI alerts HIGH RISK, and the web console catches it live via real-time IDE sync!\""
        ),
        (
            "STEP 3: The Investigation — Visual AST Diff & Risk Explainer (1:30 - 2:15)",
            "Inspect exactly what changed and understand the security implications.",
            "toolguard explain npm:dev   (or click 'Review Drift' on Web Dashboard)",
            "\"ToolGuard doesn't just throw a vague error code. It provides an automated security explanation: 'Added permission admin, network; execution command modified to curl external server'. On the web console, we get a color-coded AST visual diff showing the exact lines of code and permissions that drifted from our trusted baseline.\""
        ),
        (
            "STEP 4: The Defense — Git Pre-Commit Gate & 1-Click Restore (2:15 - 3:00)",
            "Demonstrate autonomous prevention and instantaneous recovery.",
            "git commit -m \"test\"   (Blocks commit!) -> Then: toolguard restore-test (or click 'Restore Baseline')",
            "\"Even if a developer tries to push this drifted code, ToolGuard's autonomous Git Pre-Commit Security Gate blocks the commit instantly before it ever touches CI/CD or production. To fix it, we click Restore Baseline (or run toolguard restore-test). Instantly, the workspace is clean, the baseline is verified green, and full zero-trust integrity is maintained.\""
        )
    ]

    for title, goal, action, script in steps:
        add_heading_styled(title, 2)
        p_goal = doc.add_paragraph()
        p_goal.paragraph_format.space_after = Pt(2)
        r_g1 = p_goal.add_run("Objective: ")
        r_g1.font.bold = True
        p_goal.add_run(goal)

        p_act = doc.add_paragraph()
        p_act.paragraph_format.space_after = Pt(4)
        r_a1 = p_act.add_run("Action / Command: ")
        r_a1.font.bold = True
        r_a2 = p_act.add_run(action)
        r_a2.font.name = "Consolas"
        r_a2.font.color.rgb = RGBColor(15, 23, 42)

        p_scr = doc.add_paragraph()
        p_scr.paragraph_format.space_after = Pt(10)
        r_s1 = p_scr.add_run("What to Say (Presenter Script): ")
        r_s1.font.bold = True
        r_s2 = p_scr.add_run(script)
        r_s2.font.italic = True

    # Section 4: Comparison Table
    add_heading_styled("4. Comparison: ToolGuard vs. Traditional Security Scanners", 1)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Feature / Capability"
    hdr_cells[1].text = "Traditional Tools (npm audit, Snyk)"
    hdr_cells[2].text = "ToolGuard Zero-Trust"

    for cell in hdr_cells:
        cell.paragraphs[0].runs[0].font.bold = True
        shading_xml = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls('w')))
        cell._tc.get_or_add_tcPr().append(shading_xml)

    matrix_data = [
        ("Detection Method", "Known public CVE databases", "Cryptographic SHA-256 capability baselines"),
        ("Capability Drift Detection", "None (completely blind)", "Real-time (<100ms AST diff detection)"),
        ("MCP / AI Agent Monitoring", "Unsupported", "Native Model Context Protocol tool support"),
        ("Autonomous Prevention", "Post-build CI warning only", "Pre-commit Git gate + live IDE block"),
        ("Network & Privacy Model", "Requires sending code to cloud", "100% Local / Private Network Access (zero cloud leaks)")
    ]

    for row_data in matrix_data:
        row_cells = table.add_row().cells
        for i, text in enumerate(row_data):
            row_cells[i].text = text
            row_cells[i].paragraphs[0].paragraph_format.space_after = Pt(3)

    # Section 5: Anticipated Judge Q&A
    add_heading_styled("5. Frequently Asked Judge Questions & Winning Answers", 1)

    qa_pairs = [
        (
            "Q: Why not just use git diff to see changes?",
            "A: Git diff sees raw text, not capability semantics. If an attacker reorganizes a JSON file or changes formatting, git diff shows massive noise. ToolGuard performs canonical AST normalization and zero-trust risk classification, specifically alerting on permission escalation (e.g. admin rights, new network domains, shell escapes) rather than formatting changes."
        ),
        (
            "Q: How does the web dashboard communicate with local IDEs?",
            "A: ToolGuard runs a zero-dependency native Node.js HTTP bridge on 127.0.0.1:3154 with Chromium Private Network Access (PNA) headers. The web app communicates directly with your local machine without sending any code or credentials to external servers."
        ),
        (
            "Q: What ecosystems does ToolGuard support today?",
            "A: ToolGuard features universal discovery adapters for npm package.json scripts, MCP servers (Cursor AI & Claude Desktop), VS Code tasks.json, Python scripts, Makefiles, and custom tool manifests."
        )
    ]

    for q, a in qa_pairs:
        qp = doc.add_paragraph()
        qp.paragraph_format.space_before = Pt(6)
        qp.paragraph_format.space_after = Pt(2)
        qr = qp.add_run(q)
        qr.font.bold = True
        qr.font.color.rgb = RGBColor(15, 23, 42)

        ap = doc.add_paragraph()
        ap.paragraph_format.space_after = Pt(8)
        ap.add_run(a)

    # Targets to save
    targets = [
        "ToolGuard_Live_Demo_Guide.docx",
        os.path.join("apps", "web", "public", "ToolGuard_Live_Demo_Guide.docx"),
        os.path.join("apps", "web", "dist", "ToolGuard_Live_Demo_Guide.docx")
    ]

    for target in targets:
        os.makedirs(os.path.dirname(os.path.abspath(target)), exist_ok=True)
        doc.save(target)
        print(f"[OK] Saved: {target}")

if __name__ == "__main__":
    create_demo_document()
